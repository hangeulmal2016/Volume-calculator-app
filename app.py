import streamlit as st
import pandas as pd
import numpy as np
from scipy.spatial import Delaunay
from scipy.interpolate import LinearNDInterpolator
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import ezdxf

# 1. Cấu hình giao diện Web App
st.set_page_config(page_title="Tính Toán Thể Tích Lưới Ô Vuông (TIN)", layout="wide")
st.title("📊 Ứng Dụng Tính Thể Tích Địa Hình Qua Lưới Ô Vuông (TIN)")
st.caption("Ứng dụng tối ưu hóa cho môi trường PaaS - Xử lý Stateless hoàn toàn trên bộ nhớ RAM")

# 2. Hàm đọc và xử lý file tọa độ đầu vào
def parse_coordinate_file(uploaded_file):
    if uploaded_file is None:
        return None
    try:
        bytes_data = uploaded_file.read()
        string_data = bytes_data.decode("utf-8")
        lines = [line.strip() for line in string_data.split('\n') if line.strip()]
        
        data = []
        for line in lines:
            for delim in [',', '\t', ' ']:
                parts = [p.strip() for p in line.split(delim) if p.strip()]
                if len(parts) >= 3:
                    try:
                        x = float(parts[0])
                        y = float(parts[1])
                        z = float(parts[2])
                        data.append([x, y, z])
                        break
                    except (ValueError, IndexError):
                        continue
        if len(data) == 0:
            return None
        return pd.DataFrame(data, columns=['X', 'Y', 'Z'])
    except Exception as e:
        st.error(f"Lỗi khi đọc file {uploaded_file.name}: {str(e)}")
        return None

# 3. Hàm kiểm tra điểm hình học nằm trong hay ngoài đường ranh giới kín
def is_inside_polygon(x, y, poly):
    n = len(poly)
    inside = False
    p1x, p1y = poly[0]
    for i in range(n + 1):
        p2x, p2y = poly[i % n]
        if y > min(p1y, p2y):
            if y <= max(p1y, p2y):
                if x <= max(p1x, p2x):
                    if p1y != p2y:
                        xints = (y - p1y) * (p2x - p1x) / (p2y - p1y) + p1x
                    if p1x == p2x or x <= xints:
                        inside = not inside
        p1x, p1y = p2x, p2y
    return inside

# 4. Khu vực tải tệp dữ liệu lên ứng dụng
col1, col2, col3 = st.columns(3)
with col1:
    file_bm1 = st.file_uploader("📂 Tải lên Bề mặt 1 (Hiện trạng)", type=['txt', 'csv', 'xyz'])
with col2:
    file_bm2 = st.file_uploader("📂 Tải lên Bề mặt 2 (Thiết kế / Đào đắp)", type=['txt', 'csv', 'xyz'])
with col3:
    file_boundary = st.file_uploader("📂 Tải lên File Ranh giới (Tùy chọn)", type=['txt', 'csv', 'xyz'])

# Cấu hình kích thước ô lưới vuông
grid_size = st.number_input("📐 Nhập kích thước ô lưới vuông (mét):", min_value=1.0, max_value=100.0, value=10.0, step=1.0)

# 5. Xử lý tính toán chính khi có đủ dữ liệu
if file_bm1 and file_bm2:
    df_bm1 = parse_coordinate_file(file_bm1)
    df_bm2 = parse_coordinate_file(file_bm2)
    df_bound = parse_coordinate_file(file_boundary) if file_boundary else None
    
    if df_bm1 is not None and df_bm2 is not None:
        st.success("✔️ Đã đọc thành công dữ liệu các bề mặt!")
        
        # Thiết lập ranh giới tính toán
        if df_bound is not None:
            bound_points = df_bound[['X', 'Y']].values
        else:
            combined_pts = np.vstack([df_bm1[['X', 'Y']].values, df_bm2[['X', 'Y']].values])
            min_x, min_y = combined_pts.min(axis=0)
            max_x, max_y = combined_pts.max(axis=0)
            bound_points = np.array([[min_x, min_y], [max_x, min_y], [max_x, max_y], [min_x, max_y], [min_x, min_y]])
            st.info("⚠️ Không có file ranh giới. Hệ thống tự động lấy khung bao ngoài lớn nhất làm ranh giới.")

        # Xây dựng mô hình hình học tam giác TIN
        tri_bm1 = Delaunay(df_bm1[['X', 'Y']].values)
        interp_bm1 = LinearNDInterpolator(tri_bm1, df_bm1['Z'].values)
        
        tri_bm2 = Delaunay(df_bm2[['X', 'Y']].values)
        interp_bm2 = LinearNDInterpolator(tri_bm2, df_bm2['Z'].values)
        
        # Tạo ma trận tọa độ lưới vuông
        min_x_b, min_y_b = bound_points[:, 0].min(), bound_points[:, 1].min()
        max_x_b, max_y_b = bound_points[:, 0].max(), bound_points[:, 1].max()
        
        x_coords = np.arange(min_x_b + grid_size/2, max_x_b, grid_size)
        y_coords = np.arange(min_y_b + grid_size/2, max_y_b, grid_size)
        
        grid_data = []
        grid_idx = 1
        total_cut, total_fill, total_area = 0.0, 0.0, 0.0
        
        # Vòng lặp nội suy cao độ và tính toán thể tích cho từng ô lưới
        for x in x_coords:
            for y in y_coords:
                if is_inside_polygon(x, y, bound_points):
                    z1 = float(interp_bm1(x, y))
                    z2 = float(interp_bm2(x, y))
                    
                    if np.isnan(z1) or np.isnan(z2):
                        continue
                    
                    delta_z = z2 - z1  # Dương: Đắp, Âm: Đào
                    area = grid_size * grid_size
                    volume = delta_z * area
                    
                    v_cut = abs(volume) if delta_z < 0 else 0.0
                    v_fill = volume if delta_z > 0 else 0.0
                    status = "Đào" if delta_z < 0 else ("Đắp" if delta_z > 0 else "Bằng phẳng")
                    
                    grid_data.append({
                        "ID": f"G-{grid_idx:03d}",
                        "X": round(x, 2),
                        "Y": round(y, 2),
                        "Z_BM1": round(z1, 2),
                        "Z_BM2": round(z2, 2),
                        "Delta_Z": round(delta_z, 2),
                        "Area": round(area, 2),
                        "Cut": round(v_cut, 2),
                        "Fill": round(v_fill, 2),
                        "Status": status
                    })
                    
                    total_cut += v_cut
                    total_fill += v_fill
                    total_area += area
                    grid_idx += 1
                    
        df_results = pd.DataFrame(grid_data)
        
        if df_results.empty:
            st.warning("⚠️ Không tìm thấy ô lưới nào khớp đồng thời phạm vi 2 bề mặt và ranh giới.")
        else:
            # 6. Hiển thị bảng tổng hợp KPI khối lượng
            st_col1, st_col2, st_col3, st_col4 = st.columns(4)
            st_col1.metric("Tổng diện tích tính toán", f"{total_area:,.2f} m²")
            st_col2.metric("🟥 Tổng thể tích ĐÀO (Cut)", f"{total_cut:,.2f} m³")
            st_col3.metric("🟦 Tổng thể tích ĐẮP (Fill)", f"{total_fill:,.2f} m³")
            net_vol = total_fill - total_cut
            st_col4.metric("⚖️ Thể tích THUẦN (Net)", f"{net_vol:,.2f} m³", delta=f"{'Cần Đắp' if net_vol > 0 else 'Dư Đào'}")
            
            # Hiển thị bảng số liệu trực quan
            st.subheader("📋 Bảng số liệu chi tiết ô lưới")
            st.dataframe(df_results, use_container_width=True, height=250)
            
            # 7. Ghi dữ liệu sang Excel
            excel_buffer = io.BytesIO()
            with pd.ExcelWriter(excel_buffer, engine='openpyxl') as writer:
                df_summary = pd.DataFrame([
                    ["Tên tệp bề mặt 1 (Hiện trạng)", file_bm1.name, "File"],
                    ["Tên tệp bề mặt 2 (Thiết kế)", file_bm2.name, "File"],
                    ["Tên tệp ranh giới", file_boundary.name if file_boundary else "Tự động phát sinh", "File"],
                    ["Kích thước ô lưới", f"{grid_size} x {grid_size}", "m"],
                    ["Tổng số ô lưới hợp lệ", len(df_results), "Ô"],
                    ["Diện tích vùng tính toán", total_area, "m²"],
                    ["TỔNG THỂ TÍCH ĐÀO (CUT)", total_cut, "m³"],
                    ["TỔNG THỂ TÍCH ĐẮP (FILL)", total_fill, "m³"],
                    ["THỂ TÍCH KHỐI LƯỢNG THUẦN", net_vol, "m³"]
                ], columns=["Hạng mục", "Giá trị / Thông số", "Đơn vị"])
                df_summary.to_excel(writer, sheet_name="Tổng hợp khối lượng", index=False)
                df_results.to_excel(writer, sheet_name="Chi tiết ô lưới", index=False)
                
            excel_buffer.seek(0)
            workbook = openpyxl.load_workbook(excel_buffer)
            
            # Trang trí bảng Excel
            ws1 = workbook["Tổng hợp khối lượng"]
            ws1.insert_rows(1, 2)
            ws1["A1"] = "BÁO CÁO TỔNG HỢP KHỐI LƯỢNG SAN LẤP ĐỊA HÌNH"
            ws1["A1"].font = Font(name="Arial", size=14, bold=True, color="1F4E78")
            for col in range(1, 4):
                ws1.cell(row=3, column=col).font = Font(name="Arial", bold=True)
                ws1.cell(row=3, column=col).fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            
            ws2 = workbook["Chi tiết ô lưới"]
            for col in range(1, 11):
                cell = ws2.cell(row=1, column=col)
                cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                cell.font = Font(name="Arial", color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center")
                
            for row in range(2, len(df_results) + 2):
                status_cell = ws2.cell(row=row, column=10)
                if status_cell.value == "Đào":
                    status_cell.fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
                elif status_cell.value == "Đắp":
                    status_cell.fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
                    
            styled_excel_buffer = io.BytesIO()
            workbook.save(styled_excel_buffer)
            styled_excel_buffer.seek(0)
            
            # 8. Ghi cấu trúc bản vẽ AutoCAD DXF
            doc = ezdxf.new('R2010')
            doc.layers.new(name='GRID_LINE', docproperties={'color': 8})
            doc.layers.new(name='TEXT_ID', docproperties={'color': 7})
            doc.layers.new(name='TEXT_VOLUME', docproperties={'color': 1})
            doc.layers.new(name='BOUNDARY', docproperties={'color': 3})
            msp = doc.modelspace()
            
            # Vẽ đường ranh giới bao
            for i in range(len(bound_points) - 1):
                msp.add_line(tuple(bound_points[i]), tuple(bound_points[i+1]), dxfattribs={'layer': 'BOUNDARY'})
            msp.add_line(tuple(bound_points[-1]), tuple(bound_points[0]), dxfattribs={'layer': 'BOUNDARY'})
            
            # Vẽ ô lưới và ghi chú chữ
            half_g = grid_size / 2
            for _, row_data in df_results.iterrows():
                x_c, y_c = row_data['X'], row_data['Y']
                pts = [
                    (x_c - half_g, y_c - half_g),
                    (x_c + half_g, y_c - half_g),
                    (x_c + half_g, y_c + half_g),
                    (x_c - half_g, y_c + half_g),
                    (x_c - half_g, y_c - half_g)
                ]
                for i in range(4):
                    msp.add_line(pts[i], pts[i+1], dxfattribs={'layer': 'GRID_LINE'})
                
                msp.add_text(
                    row_data['ID'],
                    dxfattribs={'layer': 'TEXT_ID', 'height': grid_size * 0.1}
                ).set_placement((x_c - half_g + 0.5, y_c + half_g - (grid_size * 0.15)))
                
                vol_color = 1 if row_data['Status'] == "Đào" else 5
                vol_str = f"C: {row_data['Cut']}" if row_data['Status'] == "Đào" else f"F: {row_data['Fill']}"
                
                msp.add_text(
                    vol_str,
                    dxfattribs={'layer': 'TEXT_VOLUME', 'height': grid_size * 0.12, 'color': vol_color}
                ).set_placement((x_c - half_g + 0.5, y_c - half_g + 0.5))
                
            dxf_buffer = io.StringIO()
            doc.write(dxf_buffer)
            dxf_bytes = dxf_buffer.getvalue().encode('utf-8')
            
            # 9. Giao diện nút bấm tải tệp kết quả
            st.subheader("📥 Tải Kết Quả Xuất File")
            dl_col1, dl_col2 = st.columns(2)
            with dl_col1:
                st.download_button(
                    label="🟢 Tải file báo cáo khối lượng Excel (.xlsx)",
                    data=styled_excel_buffer,
                    file_name="Bao_Cao_Khoi_Luong_San_Lap.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            with dl_col2:
                st.download_button(
                    label="🔵 Tải bản vẽ AutoCAD GRID (.dxf)",
                    data=dxf_bytes,
                    file_name="Ban_Ve_Luoi_O_Vuong_San_Lap.dxf",
                    mime="application/dxf"
                )
else:
    st.info("💡 Vui lòng tải lên đầy đủ hai file dữ liệu Bề mặt 1 và Bề mặt 2 ở phía trên để kích hoạt tiến trình tính toán khối lượng.")
